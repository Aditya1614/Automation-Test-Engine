import openpyxl

wb = openpyxl.load_workbook("sample_test_cases.xlsx", data_only=True)
sheet = wb.active

for r in range(2, sheet.max_row + 1):
    tc_id = sheet.cell(row=r, column=2).value
    scenario = sheet.cell(row=r, column=6).value
    title = sheet.cell(row=r, column=3).value
    flow = sheet.cell(row=r, column=1).value
    if tc_id or title:
        print(f"Row {r} | Flow: {flow} | ID: {tc_id} | Title: {title} | Scenario: {scenario}")

